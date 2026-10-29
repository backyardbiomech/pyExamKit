"""
outputs.py

The layout of ExamScanner_outputs/, and the two files in it a person uses:
gradebook.xlsx and canvas_upload.csv. docs/outputs.md describes them for
users; docs/dev/outputs-cleanup.md records why they look the way they do.

The app's working files live in app_data/: one graded results file per key
(results.csv, or results_versionA.csv and so on), each with a points file and
a grading record written by grade_functions.gradeResults. write() builds the
user's files from every results file listed in app_data/outputs.json, so a
multi-version stack produces one gradebook and one Canvas file, and a
re-grade of any version rebuilds both.
"""
from __future__ import annotations

import csv
import json
import math
import re
from dataclasses import dataclass, field, replace
from pathlib import Path

import pandas as pd

import analysis

APP_DATA = 'app_data'
GRADEBOOK = 'gradebook.xlsx'
CANVAS = 'canvas_upload.csv'
MANIFEST = 'outputs.json'

# Columns of a results file that describe the student rather than a question
ID_COLS = ('LastName', 'FirstName', 'studentID', 'form')
SCORE_COLS = ('score', 'partialscore')
NOT_STUDENTS = ('0', 'numb_correct')

DEFAULT_ID_PREFIX = 'L'
_PREFIX_SETTING = 'canvas_id_prefix'


# ── Paths ────────────────────────────────────────────────────────────────────

def app_data(outdir) -> Path:
    d = Path(outdir) / APP_DATA
    d.mkdir(parents=True, exist_ok=True)
    return d


def results_csv(outdir, version: str = '') -> str:
    """The working results file for one key; version '' for a single key."""
    name = f'results_version{version}.csv' if version else 'results.csv'
    return str(app_data(outdir) / name)


def artifact_dir(csv_path) -> Path:
    """Where a results file's companions live. Results files are in app_data/
    now; an older run kept results.csv at the top with app_data/ beside it."""
    parent = Path(csv_path).parent
    d = parent if parent.name == APP_DATA else parent / APP_DATA
    d.mkdir(parents=True, exist_ok=True)
    return d


def outdir_of(csv_path) -> Path:
    parent = Path(csv_path).parent
    return parent.parent if parent.name == APP_DATA else parent


def points_csv(csv_path) -> Path:
    return artifact_dir(csv_path) / (Path(csv_path).stem + '_points.csv')


def grading_json(csv_path) -> Path:
    return artifact_dir(csv_path) / (Path(csv_path).stem + '_grading.json')


def version_of(csv_path) -> str:
    m = re.fullmatch(r'results_version([A-Z])', Path(csv_path).stem)
    return m[1] if m else ''


def record(outdir, csv_paths) -> None:
    """Name the results files this scan produced, so write() leaves out any a
    different, earlier scan into the same folder left behind."""
    names = sorted(Path(p).name for p in csv_paths)
    (app_data(outdir) / MANIFEST).write_text(json.dumps({'results': names}, indent=2),
                                             encoding='utf-8')
    import edits
    moved = edits.set_aside(outdir)
    if moved:
        print(f'[Outputs] Edits made to the earlier scan here were set aside → {moved.name}',
              flush=True)


def listed(outdir) -> list[Path]:
    """The results files write() combines: those recorded by the last scan,
    else every results file in app_data/."""
    d = Path(outdir) / APP_DATA
    try:
        names = json.loads((d / MANIFEST).read_text(encoding='utf-8'))['results']
        return [d / n for n in names if (d / n).exists()]
    except (OSError, ValueError, KeyError):
        return sorted(p for p in d.glob('results*.csv')
                      if re.fullmatch(r'results(_version[A-Z])?', p.stem))


def id_prefix() -> str:
    """What goes in front of a student ID in the Canvas file ('L' at Longwood)."""
    import ai_ocr
    return str(ai_ocr.load_config().get(_PREFIX_SETTING, DEFAULT_ID_PREFIX))


def set_id_prefix(prefix: str) -> None:
    import ai_ocr
    ai_ocr.save_config({_PREFIX_SETTING: prefix.strip()})


# ── One graded results file ─────────────────────────────────────────────────

@dataclass
class Graded:
    version: str
    df: pd.DataFrame            # answers; row '0' is the key
    pts: pd.DataFrame           # points earned, same shape
    possible: dict              # {question column: points}
    title: str = ''
    sources: dict = field(default_factory=dict)
    choices: dict = field(default_factory=dict)
    texts: dict = field(default_factory=dict)            # {question column: question text}
    open_answers: dict = field(default_factory=dict)   # {openQ_N: [accepted answers]}
    settings: dict = field(default_factory=dict)       # how it was scored (grading record)
    accepted: dict = field(default_factory=dict)       # {column: [keys]} after a key edit
    dropped: set = field(default_factory=set)          # columns a credit edit dropped

    def key_label(self, col: str, bank: bool = False) -> str:
        """The key as shown: every accepted answer, in the bank's letters
        when bank is set."""
        keys = self.accepted.get(col) or [str(self.df.loc['0', col])]
        return ' or '.join(self.bank_letters(col, k) if bank else k for k in keys)

    @property
    def students(self) -> list[str]:
        return [r for r in self.df.index if r not in NOT_STUDENTS]

    @property
    def q_cols(self) -> list[str]:
        return [c for c in self.df.columns if c not in ID_COLS and c not in SCORE_COLS]

    @property
    def id_cols(self) -> list[str]:
        return [c for c in ID_COLS if c in self.df.columns]

    def total(self, row: str) -> float:
        return sum(self.earned(row, c) for c in self.possible)

    def earned(self, row: str, col: str) -> float:
        try:
            v = float(self.pts.loc[row, col])
        except (KeyError, TypeError, ValueError):
            return 0.0
        return 0.0 if math.isnan(v) else v

    def asked(self, row: str, col: str) -> bool:
        """False for a question not on this student's sheet (a practical form)."""
        return not pd.isna(self.df.loc[row, col])

    def bank_letters(self, col: str, letters: str) -> str:
        """A bubbled answer in the bank's lettering, through the key's choices."""
        order = self.choices.get(col, '')
        if not order:
            return letters
        out = []
        for ch in letters:
            i = ord(ch) - ord('A')
            out.append(order[i] if 0 <= i < len(order) else ch)
        return ''.join(sorted(out))

    def sheet_letters(self, col: str, letters: str) -> str:
        """The reverse of bank_letters: bank letters as they sat on this sheet.
        A letter the question did not have is left as it is."""
        order = self.choices.get(col, '')
        if not order:
            return letters
        return ''.join(sorted(chr(ord('A') + order.index(ch)) if ch in order else ch
                              for ch in letters))

    def points_possible(self) -> float:
        return sum(self.possible.values())


def _read_indexed(path) -> pd.DataFrame:
    df = pd.read_csv(path, dtype=object)
    first = df.columns[0]
    df = df.set_index(first)
    df.index = df.index.map(str)
    df.index.name = None
    return df


def load(csv_path) -> Graded:
    df = _read_indexed(csv_path)
    pts_path = points_csv(csv_path)
    pts = _read_indexed(pts_path) if pts_path.exists() else df.copy()
    rec = {}
    try:
        rec = json.loads(grading_json(csv_path).read_text(encoding='utf-8'))
    except (OSError, ValueError):
        pass
    open_answers = {}
    try:
        raw = json.loads((artifact_dir(csv_path) / (Path(csv_path).stem + '_openq_answers.json'))
                         .read_text(encoding='utf-8'))
        for qk, v in raw.items():
            open_answers[qk] = v.get('full', []) if isinstance(v, dict) else list(v)
    except (OSError, ValueError, AttributeError):
        pass
    return Graded(version=version_of(csv_path), df=df, pts=pts,
                  possible=rec.get('possible', {}), title=rec.get('title', ''),
                  sources=rec.get('sources', {}), choices=rec.get('choices', {}),
                  texts=rec.get('texts', {}), open_answers=open_answers,
                  settings=rec.get('settings', {}))


# ── Writing ──────────────────────────────────────────────────────────────────

def write(outdir, prefix: str | None = None, csvs=None) -> bool:
    """Write gradebook.xlsx and canvas_upload.csv from the listed results,
    or from csvs when given. True when both were saved; then the edits they
    carry are recorded as exported."""
    graded = [load(p) for p in (csvs if csvs is not None else listed(outdir))]
    graded = [g for g in graded if g.students]
    if not graded:
        print('[Outputs] No graded results to write.', flush=True)
        return False
    import edits
    made = edits.load(outdir)
    graded, problems = edits.apply(graded, made)
    for msg in problems:
        print(f'[Outputs] {msg}', flush=True)
    adjust = edits.adjustments(made)
    title = next((g.title for g in graded if g.title), '') or Path(outdir).resolve().parent.name
    prefix = id_prefix() if prefix is None else prefix
    saved = True
    # Either file fails while it is open in Excel on Windows, which locks it
    for name, fn in ((GRADEBOOK, write_gradebook), (CANVAS, write_canvas)):
        try:
            fn(Path(outdir) / name, graded, title, prefix, adjust)
            print(f'{"Gradebook" if name == GRADEBOOK else "Canvas file"} saved → '
                  f'{Path(outdir) / name}', flush=True)
        except Exception as exc:
            saved = False
            print(f'[Outputs] Could not save {name}: {exc}', flush=True)
    if saved and csvs is None:
        edits.mark_exported(outdir, made)
    return saved


def _score(x: float) -> str:
    return f'{round(x, 2):g}'


def _canvas_rows(graded: list[Graded], prefix: str) -> list[tuple]:
    """The Canvas file's students in its order, sorted by name, as
    (Student, SIS User ID, index into graded, results row). The gradebook's
    Canvas scores tab uses the same order, so it pastes over the file."""
    rows = []
    for i, g in enumerate(graded):
        for r in g.students:
            last, first, sid = (str(g.df.loc[r, c]) if not pd.isna(g.df.loc[r, c]) else ''
                                for c in ('LastName', 'FirstName', 'studentID'))
            # Name bubbles pad with dashes, and v2 sheets have none to read
            last, first = (x.rstrip('-') for x in (last, first))
            name = f'{last}, {first}' if first else last
            rows.append((last.lower(), first.lower(), sid, name,
                         f'{prefix}{sid}' if sid else '', i, r))
    rows.sort(key=lambda row: row[:3])
    return [row[3:] for row in rows]


def adjustment(adjust: dict, version: str) -> float:
    """Points added to a version's students: the class's plus the version's."""
    return (adjust or {}).get('', 0) + ((adjust or {}).get(version, 0) if version else 0)


def write_canvas(path, graded: list[Graded], title: str, prefix: str,
                 adjust: dict | None = None) -> None:
    """A file Canvas's gradebook import takes as is: its own header, a Points
    Possible row, then one row per student, sorted by name. Canvas matches
    students on SIS User ID and asks which assignment the column is. The
    scores carry the adjustments ({version: points}, '' the class)."""
    possible = {g.points_possible() for g in graded}
    if len(possible) > 1:
        print('[Outputs] The versions are worth different totals ('
              + ', '.join(_score(p) for p in sorted(possible))
              + '); Points Possible uses the largest.', flush=True)
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['Student', 'SIS User ID', title])
        w.writerow(['    Points Possible', '', _score(max(possible))])
        for name, sis, i, r in _canvas_rows(graded, prefix):
            w.writerow([name, sis, _score(graded[i].total(r)
                                          + adjustment(adjust, graded[i].version))])


def write_gradebook(path, graded: list[Graded], title: str = '', prefix: str = '',
                    adjust: dict | None = None) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    stats = wb.create_sheet('Exam stats')
    cells = _adjust_cells([g.version for g in graded if g.students])
    by_version = []
    for g in graded:
        name = f'Version {g.version}' if g.version else 'Gradebook'
        refs = [cells['']] + ([cells[g.version]] if g.version in cells else [])
        by_version.append(_gradebook_sheet(wb.create_sheet(name), g, refs))
    ranges = [r for r in by_version if r]
    items = item_table(graded)
    combined = None
    if len(graded) > 1 and items.by_source:
        combined = _by_question_sheet(wb.create_sheet('By question'), graded, items)
    if len(ranges) == 1:
        combined = ranges[0]
    _stats_sheet(stats, ranges, combined, next((g.title for g in graded if g.title), ''),
                 adjust or {})
    _item_sheet(wb.create_sheet('Item analysis'), items)
    _canvas_sheet(wb.create_sheet('Canvas scores'), graded, by_version,
                  title or 'Score', prefix)
    wb.save(path)


# ── Sheets ───────────────────────────────────────────────────────────────────

def _styles():
    from openpyxl.styles import Alignment, Font, PatternFill
    return {
        'hdr_font': Font(bold=True),
        'hdr_align': Alignment(wrap_text=True, horizontal='center', vertical='center'),
        'hdr_fill': PatternFill('solid', fgColor='BDD7EE'),
        'key_fill': PatternFill('solid', fgColor='FFFF99'),
        'flag_fill': PatternFill('solid', fgColor='F8CBAD'),
    }


def _header(ws, values, height=36):
    st = _styles()
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font, cell.alignment, cell.fill = st['hdr_font'], st['hdr_align'], st['hdr_fill']
    ws.row_dimensions[ws.max_row].height = height


def _text_row(ws, first_col: int, texts: list[str]) -> None:
    """A row of question text, each spanning its question's answer and
    points columns (texts[i] over columns first_col + 2i and the next)."""
    from openpyxl.styles import Alignment
    ws.append(['Question'])
    r = ws.max_row
    for i, text in enumerate(texts):
        col = first_col + 2 * i
        cell = ws.cell(row=r, column=col)
        cell.value = text or None
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
    ws.row_dimensions[r].height = 90


# Canvas's discrimination index compares the top and bottom 27% of the class
# by total score (docs/outputs.md).
DI_GROUP = analysis.DI_GROUP


def _summary_rows(ws, first: int, last: int, pts_cols: list, total_col: int,
                  more_totals: tuple = ()) -> None:
    """Mean, median, and discrimination index under each points column, as
    formulas over student rows first..last, so they follow edited points the
    way the totals do. pts_cols is [(column number, points possible)];
    more_totals are columns like Adjusted that get a mean and median too."""
    from openpyxl.utils import get_column_letter
    st = _styles()
    tot = get_column_letter(total_col)
    t_rng = f'${tot}${first}:${tot}${last}'
    rows = {}
    for label in ('Mean', 'Median', 'Discrimination index'):
        ws.append([label])
        rows[label] = ws.max_row
        for cell in ws[ws.max_row]:
            cell.font, cell.fill = st['hdr_font'], st['key_fill']
    for col, possible in pts_cols + [(c, None) for c in (total_col, *more_totals)]:
        c = get_column_letter(col)
        rng = f'{c}{first}:{c}{last}'
        ws.cell(row=rows['Mean'], column=col).value = f'=IFERROR(ROUND(AVERAGE({rng}),2),"")'
        ws.cell(row=rows['Median'], column=col).value = f'=IFERROR(ROUND(MEDIAN({rng}),2),"")'
        if possible:
            # Mean fraction of the points earned by students at or above the
            # 73rd percentile of totals, less that of students at or below
            # the 27th; ties at a cutoff join the group
            upper = f'AVERAGEIFS({rng},{t_rng},">="&PERCENTILE({t_rng},{1 - DI_GROUP}))'
            lower = f'AVERAGEIFS({rng},{t_rng},"<="&PERCENTILE({t_rng},{DI_GROUP}))'
            ws.cell(row=rows['Discrimination index'], column=col).value = (
                f'=IFERROR(ROUND(({upper}-{lower})/{possible},2),"")')


def _gradebook_sheet(ws, g: Graded, adjust: list[str] = ()):
    """Row 1 headers, row 2 the key (yellow), then a row per student: each
    question's answer and points, and a Total that is a live SUM of the
    points, so correcting a points cell after a regrade updates the total.
    Adjusted is the total plus the adjustment cells on Exam stats."""
    from openpyxl.utils import get_column_letter
    st = _styles()
    ids, qs = g.id_cols, g.q_cols
    n_id = len(ids)

    header = list(ids)
    for qc in qs:
        key = str(g.df.loc['0', qc])
        if qc in g.dropped:
            label = f'{qc}\n(dropped)'
        elif key == 'CC':
            label = f'{qc}\n(written)'
        elif key in ('ignore', 'nan', ''):
            label = f'{qc}\n(not graded)'
        else:
            label = f'{qc}\n(Key: {g.key_label(qc)})'
        header += [label, f'{qc} Pts']
    header += ['Total', 'Adjusted']
    _header(ws, header)
    has_text = any(g.texts.get(qc) for qc in qs)
    if has_text:
        _text_row(ws, n_id + 1, [g.texts.get(qc, '') for qc in qs])

    key_row = ['KEY'] + [''] * (n_id - 1)
    for qc in qs:
        key = str(g.df.loc['0', qc])
        if key == 'CC' and g.open_answers.get(qc):
            key = ' | '.join(g.open_answers[qc])
        elif qc in g.accepted:
            key = g.key_label(qc)
        key_row += ['' if key == 'nan' else key, g.possible.get(qc, '')]
    key_row.append(g.points_possible())
    ws.append(key_row)
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = st['hdr_font'], st['key_fill']

    first = ws.max_row + 1
    for r in g.students:
        row = [('' if pd.isna(g.df.loc[r, c]) else str(g.df.loc[r, c])) for c in ids]
        for qc in qs:
            ans = g.df.loc[r, qc]
            # Blank, not 0, for a question not on this student's form, so
            # the question's mean and median leave the student out
            row += ['' if pd.isna(ans) else str(ans),
                    g.earned(r, qc) if g.asked(r, qc) else None]
        ws.append(row)
        refs = ','.join(f'{get_column_letter(n_id + 2 + 2 * i)}{ws.max_row}'
                        for i in range(len(qs)))
        ws.cell(row=ws.max_row, column=n_id + 1 + 2 * len(qs)).value = (
            f'=SUM({refs})' if refs else 0)
        total = f'{get_column_letter(n_id + 1 + 2 * len(qs))}{ws.max_row}'
        ws.cell(row=ws.max_row, column=n_id + 2 + 2 * len(qs)).value = (
            '=' + '+'.join([total] + [f'N({a})' for a in adjust]))
    if g.students:
        _summary_rows(ws, first, ws.max_row,
                      [(n_id + 2 + 2 * i, g.possible.get(qc, 0)) for i, qc in enumerate(qs)],
                      n_id + 1 + 2 * len(qs), (n_id + 2 + 2 * len(qs),))

    ws.freeze_panes = ws.cell(row=first, column=n_id + 1)
    for i, w in enumerate([16, 14, 12, 8][:n_id], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i in range(len(qs)):
        ws.column_dimensions[get_column_letter(n_id + 1 + 2 * i)].width = 18
        ws.column_dimensions[get_column_letter(n_id + 2 + 2 * i)].width = 8
    ws.column_dimensions[get_column_letter(n_id + 1 + 2 * len(qs))].width = 10
    ws.column_dimensions[get_column_letter(n_id + 2 + 2 * len(qs))].width = 10
    if not g.students:
        return None
    # Where the totals are, for the Exam stats and Canvas scores tabs
    return TotalRange(g.version, ws.title, get_column_letter(n_id + 1 + 2 * len(qs)),
                      first, first + len(g.students) - 1, g.points_possible(),
                      [get_column_letter(n_id + 2 + 2 * i) for i, qc in enumerate(qs)
                       if g.possible.get(qc, 0) > 0],
                      get_column_letter(n_id + 2 + 2 * len(qs)))


# ── Exam stats ───────────────────────────────────────────────────────────────

@dataclass
class TotalRange:
    version: str
    sheet: str
    col: str
    first: int
    last: int
    possible: float
    items: list = field(default_factory=list)    # each graded question's points column
    adj_col: str = ''                            # the Adjusted column, if any

    def adjusted(self) -> TotalRange:
        """The same students' adjusted totals."""
        return replace(self, col=self.adj_col)

    def col_ref(self, col: str) -> str:
        return f"'{self.sheet}'!${col}${self.first}:${col}${self.last}"

    @property
    def ref(self) -> str:
        return self.col_ref(self.col)

    def alpha(self) -> str | None:
        """Cronbach's alpha as a formula: k/(k-1) x (1 - the items' summed
        variances over the totals' variance). On all-or-nothing questions it
        is KR-20."""
        k = len(self.items)
        if k < 2:
            return None
        items = ','.join(f'VARP({self.col_ref(c)})' for c in self.items)
        return f'=IFERROR(ROUND({k}/{k - 1}*(1-SUM({items})/VARP({self.ref})),2),"")'


# Letter-grade bands, shared with the Results tab (analysis.py)
BANDS = analysis.BANDS
BAR_COLOR = '2A78D6'
UNADJUSTED_COLOR = 'B4B2AC'


# The adjustment inputs sit at the top of Exam stats, a fixed place the
# gradebook tabs' Adjusted formulas can point at before the tab is written
ADJUST_ROW = 4


def _adjust_cells(versions: list[str]) -> dict:
    """Where each adjustment is typed: '' for the whole class, and one per
    version when there are several. A version's students get both."""
    from openpyxl.utils import get_column_letter
    cells = {'': f"'Exam stats'!$B${ADJUST_ROW}"}
    if len(versions) > 1:
        for i, v in enumerate(versions):
            cells[v] = f"'Exam stats'!${get_column_letter(3 + i)}${ADJUST_ROW}"
    return cells


def _stats_sheet(ws, ranges: list[TotalRange], combined: TotalRange | None,
                 title: str, adjust: dict | None = None) -> None:
    """The adjustment inputs, then summary statistics of the totals for the
    class and for each version, each beside its adjusted figures, and how
    many students earn each letter grade before and after adjusting. Every
    number is a formula over the gradebook tabs, so it follows corrected
    points and a changed adjustment. combined is the whole class question
    by question (By question, or the one version), which the class's
    reliability needs; None leaves it blank. adjust is the Results tab's
    adjustments, the Points added row's starting values."""
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    st = _styles()
    note_font = Font(italic=True, color='52514E')
    ws.append([title or 'Exam stats'])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([])
    if not ranges:
        ws.append(['No graded students.'])
        return

    groups = [('All students', ranges)]
    if len(ranges) > 1:
        groups += [(f'Version {r.version}', [r]) for r in ranges]

    # Adjustment: typed here, added to every total on the gradebook tabs
    _header(ws, ['Adjustment'] + [label for label, _ in groups], height=20)
    adjust = adjust or {}
    if len(groups) > 1:
        added = [adjust.get('', 0)] + [adjust.get(r.version, 0) for r in ranges]
    else:
        added = [adjustment(adjust, ranges[0].version)]
    ws.append(['Points added'] + [_num(a) for a in added])
    assert ws.max_row == ADJUST_ROW
    for cell in ws[ADJUST_ROW][1:]:
        cell.fill = st['key_fill']
        cell.font = Font(bold=True, color='0000FF')
    ws.append(['Type the points to add to every score in the yellow cells. '
               + ('A student gets the class amount plus their version\'s. '
                  if len(groups) > 1 else '')
               + 'The Adjusted columns here, on the gradebook tabs, and on Canvas '
                 'scores follow.'])
    ws.cell(row=ws.max_row, column=1).font = note_font
    ws.append([])

    # Scores: each group's totals, then its adjusted totals beside them
    cols = []
    for label, rs in groups:
        cols += [(label, rs, False), ('Adjusted', [r.adjusted() for r in rs], True)]
    _header(ws, ['Scores'] + [label for label, _, _ in cols], height=20)
    same_possible = len({r.possible for r in ranges}) == 1

    def union(rs):
        return ','.join(r.ref for r in rs)

    stats = [('Students', 'COUNT({})', False), ('Mean', 'ROUND(AVERAGE({}),2)', True),
             ('Median', 'ROUND(MEDIAN({}),2)', True),
             ('Standard deviation', 'ROUND(STDEV({}),2)', False),
             ('Highest', 'MAX({})', True), ('Lowest', 'MIN({})', True)]
    ws.append(['Points possible'] + [_num(rs[0].possible) if same_possible or len(rs) == 1
                                     else 'varies' for _, rs, _ in cols])
    pct_rows = []
    for label, f, as_pct in stats:
        ws.append([label] + [f'=IFERROR({f.format(union(rs))},"")' for _, rs, _ in cols])
        if label == 'Standard deviation':
            # Adding points moves no student relative to another on a version,
            # so reliability is the unadjusted figure's and left blank
            alphas = []
            for i, (_, rs, adj) in enumerate(cols):
                if adj:
                    alphas.append(None)
                elif len(rs) > 1 or i == 0:
                    alphas.append(combined.alpha() if combined else None)
                else:
                    alphas.append(rs[0].alpha())
            ws.append(["Reliability (Cronbach's alpha)"] + alphas)
        if as_pct:
            row = [f'{label} (% of possible)']
            for i, (_, rs, _) in enumerate(cols):
                if same_possible or len(rs) == 1:
                    ref = f'{get_column_letter(2 + i)}{ws.max_row}'
                    row.append(f'=IFERROR({ref}/{rs[0].possible},"")')
                else:
                    row.append(None)
            ws.append(row)
            pct_rows.append(ws.max_row)
    for r in pct_rows:
        for cell in ws[r][1:]:
            cell.number_format = '0.0%'

    # Distribution across the whole class, each student against their own
    # version's points possible, before and after adjusting
    ws.append([])
    ws.append([])
    _header(ws, ['Grade', 'Percent of possible', 'Students', '% of class',
                 'Adjusted', 'Adjusted % of class'], height=20)
    first = ws.max_row + 1
    n_row = first + len(BANDS)

    def count(rs, lo, hi):
        terms = []
        for r in rs:
            conds = []
            if lo is not None:
                conds.append(f'{r.ref},">="&{lo / 100}*{r.possible}')
            if hi is not None:
                conds.append(f'{r.ref},"<"&{hi / 100}*{r.possible}')
            terms.append(f'COUNTIFS({",".join(conds)})')
        return '=' + '+'.join(terms)

    adjusted = [r.adjusted() for r in ranges]
    for grade, lo, hi, label in BANDS:
        row = ws.max_row + 1
        ws.append([grade, label, count(ranges, lo, hi), f'=IFERROR(C{row}/$C${n_row},"")',
                   count(adjusted, lo, hi), f'=IFERROR(E{row}/$E${n_row},"")'])
        for c in (4, 6):
            ws.cell(row=row, column=c).number_format = '0%'
    last = ws.max_row
    ws.append(['Total', ''] + [f'=SUM({c}{first}:{c}{last})' for c in 'CDEF'])
    for c in (4, 6):
        ws.cell(row=ws.max_row, column=c).number_format = '0%'
    for cell in ws[ws.max_row]:
        cell.font = st['hdr_font']
    ws.append(['Each band runs from its lower bound up to, not including, the next: '
               '82.9% is a B−. An adjusted score above the points possible counts as an A.'])
    ws.cell(row=ws.max_row, column=1).font = note_font

    # Unadjusted in gray behind the adjusted in blue
    chart = BarChart()
    chart.type = 'col'
    chart.title = 'Students by grade'
    chart.y_axis.title = 'Students'
    chart.legend.position = 'b'
    chart.gapWidth = 30
    for col, color, name in ((3, UNADJUSTED_COLOR, 'Unadjusted'), (5, BAR_COLOR, 'Adjusted')):
        chart.add_data(Reference(ws, min_col=col, min_row=first, max_row=last),
                       titles_from_data=False)
        series = chart.series[-1]
        series.tx = SeriesLabel(v=name)
        series.graphicalProperties = GraphicalProperties(solidFill=color)
        series.graphicalProperties.line.noFill = True
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E0E0DD'))
    chart.y_axis.number_format = '0'
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, f'{get_column_letter(max(len(cols), 6) + 3)}3')

    ws.column_dimensions['A'].width = 30
    for i in range(len(cols)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 14


def _canvas_sheet(ws, graded: list[Graded], ranges: list, title: str, prefix: str) -> None:
    """canvas_upload.csv with the adjusted scores: the same header, Points
    Possible row, and students in the same order, so column C pastes over the
    file's scores. ranges[i] locates graded[i]'s totals. The unadjusted
    scores sit to the right, past a blank column, for checking."""
    from openpyxl.styles import Font
    possible = max(g.points_possible() for g in graded)
    _header(ws, ['Student', 'SIS User ID', title, None, 'Unadjusted'], height=20)
    ws.append(['    Points Possible', '', _num(possible)])
    for name, sis, i, r in _canvas_rows(graded, prefix):
        t = ranges[i]
        row = t.first + graded[i].students.index(r)
        ws.append([name, sis, f"=ROUND('{t.sheet}'!{t.adj_col}{row},2)", None,
                   f"=ROUND('{t.sheet}'!{t.col}{row},2)"])
    ws.cell(row=1, column=7).value = (
        'Column C is canvas_upload.csv with the adjustment on Exam stats added, row for row. '
        'Copy C3 down to the last student and paste it over C3 in canvas_upload.csv.')
    ws.cell(row=1, column=7).font = Font(italic=True, color='52514E')
    ws.freeze_panes = 'A3'
    for col, w in zip('ABCDE', (28, 14, 14, 3, 12)):
        ws.column_dimensions[col].width = w


def _num(x: float):
    return int(x) if float(x).is_integer() else round(x, 2)


# ── Item analysis ────────────────────────────────────────────────────────────

@dataclass
class Item:
    label: str                       # the bank question, or the sheet's number
    where: list[str]                 # where it sat on each version's sheet
    key: str                         # in bank letters
    points: float
    written: bool
    text: str = ''
    cols: list = field(default_factory=list)     # [(version, column)], beside where
    # (student uid, fraction of points earned, points earned, answer in bank letters)
    responses: list = field(default_factory=list)


@dataclass
class ItemTable:
    items: list[Item]
    totals: dict                      # {student uid: total points}
    by_source: bool                   # items matched across versions by bank question


def _qnum(col: str) -> str:
    """A question column as its number on the sheet: Q012 is 12, openQ_1A is 1A."""
    m = re.fullmatch(r'Q0*(\d+)', col)
    if m:
        return m[1]
    return col.split('_', 1)[1] if col.startswith('openQ_') else col


def item_table(graded: list[Graded]) -> ItemTable:
    """One item per question. With every version's key carrying sources,
    versions' questions are matched by bank question; otherwise each
    version's questions stand alone."""
    by_source = all(g.sources for g in graded)
    items: dict[str, Item] = {}
    totals = {}
    for g in graded:
        for r in g.students:
            totals[(g.version, r)] = g.total(r)
        for col, pts in g.possible.items():
            if pts <= 0 or col not in g.df.columns:
                continue            # the version question, worth nothing
            key = str(g.df.loc['0', col])
            written = key == 'CC'
            if by_source:
                ident = g.sources.get(col)
                if not ident:
                    continue
            else:
                ident = f'{g.version}{_qnum(col)}' if len(graded) > 1 else _qnum(col)
            where = f'{g.version}{_qnum(col)}' if g.version else _qnum(col)
            it = items.get(ident)
            if it is None:
                it = items[ident] = Item(
                    label=ident, where=[], points=pts, written=written,
                    text=g.texts.get(col, ''),
                    key='written' if written else g.key_label(col, bank=True))
            it.where.append(where)
            it.cols.append((g.version, col))
            for r in g.students:
                if not g.asked(r, col):
                    continue
                earned = g.earned(r, col)
                ans = str(g.df.loc[r, col])
                ans = ans[:2] if written else g.bank_letters(col, ans.replace('-', ''))
                it.responses.append(((g.version, r), earned / pts, earned, ans))
    return ItemTable(list(items.values()), totals, by_source)


def _pearson(xs: list[float], ys: list[float]) -> float | None:
    n = len(xs)
    if n < 3:
        return None
    mx, my = sum(xs) / n, sum(ys) / n
    sxx = sum((x - mx) ** 2 for x in xs)
    syy = sum((y - my) ** 2 for y in ys)
    if sxx == 0 or syy == 0:
        return None
    return sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / math.sqrt(sxx * syy)


def item_stats(it: Item, totals: dict) -> dict:
    """Difficulty (mean fraction of points earned) and discrimination (the
    corrected item-total correlation: the item's score against the total
    without it), plus how often each choice was marked."""
    n = len(it.responses)
    fracs = [f for _, f, _, _ in it.responses]
    rest = [totals[uid] - earned for uid, _, earned, _ in it.responses]
    counts: dict[str, int] = {}
    for *_, ans in it.responses:
        for ch in (ans or ' ') if not it.written else [ans]:
            k = 'Blank' if ch == ' ' else ch
            counts[k] = counts.get(k, 0) + 1
    return {'n': n, 'mean': sum(fracs) / n if n else None,
            'r': _pearson(fracs, rest), 'counts': counts}


# Below this corrected item-total correlation a question is flagged for a
# look; 0.2 is a common rule of thumb, not a law (docs/outputs.md).
FLAG_BELOW = 0.2


def _item_sheet(ws, t: ItemTable) -> None:
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    st = _styles()
    letters = sorted({k for it in t.items if not it.written
                      for k in item_stats(it, t.totals)['counts'] if k != 'Blank'})
    grades = ['CC', 'CX', 'XX']
    header = ['Question', 'Text', 'On sheet', 'Key', 'Points', 'Students', 'Mean score',
              'Discrimination', 'Look at'] + letters + ['Blank'] + grades
    _header(ws, header)
    for it in t.items:
        s = item_stats(it, t.totals)
        r = s['r']
        flag = '' if r is None or r >= FLAG_BELOW else ('negative' if r < 0 else 'low')
        row = [it.label, it.text, ', '.join(it.where), it.key, it.points, s['n'],
               None if s['mean'] is None else round(s['mean'], 3),
               None if r is None else round(r, 3), flag]
        row += [s['counts'].get(k, 0) if not it.written else '' for k in letters]
        row.append(s['counts'].get('Blank', 0) if not it.written else '')
        row += [s['counts'].get(g, 0) if it.written else '' for g in grades]
        ws.append(row)
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical='top')
        if flag:
            for c in ws[ws.max_row][:9]:
                c.fill = st['flag_fill']
        # Bold the key's letters among the choice counts
        for i, k in enumerate(letters):
            if k in it.key:
                ws.cell(row=ws.max_row, column=10 + i).font = st['hdr_font']
    ws.freeze_panes = 'B2'
    for i, w in enumerate([18, 50, 14, 8, 8, 9, 10, 14, 9], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _by_question_sheet(ws, graded: list[Graded], t: ItemTable) -> TotalRange | None:
    """Every student in one table, a column pair per bank question, answers
    in the bank's lettering: the raw material for any further analysis."""
    from openpyxl.utils import get_column_letter
    header = ['LastName', 'FirstName', 'studentID', 'Version']
    for it in t.items:
        header += [f'{it.label}\n(Key: {it.key})', f'{it.label} Pts']
    header.append('Total')
    _header(ws, header)
    if any(it.text for it in t.items):
        _text_row(ws, 5, [it.text for it in t.items])
    first = ws.max_row + 1
    cells = {}
    for i, it in enumerate(t.items):
        for uid, _, earned, ans in it.responses:
            cells[(uid, i)] = (ans, earned)
    rows = []
    for g in graded:
        for r in g.students:
            uid = (g.version, r)
            row = [('' if pd.isna(g.df.loc[r, c]) else str(g.df.loc[r, c]))
                   for c in ('LastName', 'FirstName', 'studentID')] + [g.version]
            for i in range(len(t.items)):
                ans, earned = cells.get((uid, i), ('', None))
                row += [ans, earned]
            rows.append(row)
    for row in sorted(rows, key=lambda r: (r[0].lower(), r[1].lower(), r[2])):
        ws.append(row)
        refs = ','.join(f'{get_column_letter(6 + 2 * i)}{ws.max_row}'
                        for i in range(len(t.items)))
        ws.cell(row=ws.max_row, column=5 + 2 * len(t.items)).value = (
            f'=SUM({refs})' if refs else 0)
    if rows:
        _summary_rows(ws, first, ws.max_row,
                      [(6 + 2 * i, it.points) for i, it in enumerate(t.items)],
                      5 + 2 * len(t.items))
    ws.freeze_panes = ws.cell(row=first, column=5)
    for i in range(len(t.items)):
        ws.column_dimensions[get_column_letter(5 + 2 * i)].width = 14
        ws.column_dimensions[get_column_letter(6 + 2 * i)].width = 7
    if not rows:
        return None
    return TotalRange('', ws.title, get_column_letter(5 + 2 * len(t.items)), first,
                      first + len(rows) - 1, max(g.points_possible() for g in graded),
                      [get_column_letter(6 + 2 * i) for i in range(len(t.items))])
