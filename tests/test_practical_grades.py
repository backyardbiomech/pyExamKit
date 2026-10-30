'''
A lab practical's points possible. One key carries every letter at every
station, and each student answers the two their form holds, so what a student
could earn is half the key's total and the analysis has to score them against
their own sheet.

The fixture is a two-station practical, questions A to D at each station,
worth a point each, with two forms:

    AC  ABEL  2A 2C right, 1A 1C wrong      2 of 4
    BD  BIRD  1B 1D 2B right, 2D wrong      3 of 4

    python -m unittest tests.test_practical_grades -v
'''
import contextlib
import csv
import io
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'tests'))

import edits  # noqa: E402
import grade_functions  # noqa: E402
import outputs  # noqa: E402
import results_model as rm  # noqa: E402

QS = ['openQ_1A', 'openQ_1B', 'openQ_1C', 'openQ_1D',
      'openQ_2A', 'openQ_2B', 'openQ_2C', 'openQ_2D']
# What the grader marked, blank for a question not on that student's form
STUDENTS = [('ABEL', '1', {'openQ_1A': 'XX', 'openQ_1C': 'XX',
                           'openQ_2A': 'CC', 'openQ_2C': 'CC'}),
            ('BIRD', '2', {'openQ_1B': 'CC', 'openQ_1D': 'CC',
                           'openQ_2B': 'CC', 'openQ_2D': 'XX'})]


def quiet(fn, *a, **kw):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **kw)


def make_outputs(root: Path, points: dict | None = None) -> tuple[Path, Path]:
    """An outputs folder holding the practical graded, as a scan leaves it.
    points overrides what a question is worth, as a (2 pts) line does."""
    out = root / 'ExamScanner_outputs'
    (out / 'marked').mkdir(parents=True)
    rows = [['0', 'KEY', '', ''] + ['CC'] * len(QS)]
    rows += [[str(i), last, 'X', sid] + [marks.get(q, '') for q in QS]
             for i, (last, sid, marks) in enumerate(STUDENTS, 1)]
    path = Path(outputs.results_csv(out))
    pd.DataFrame(rows, columns=['index', 'LastName', 'FirstName', 'studentID'] + QS
                 ).to_csv(path, index=False)
    quiet(grade_functions.gradeResults, str(path), False, True, 1, 1, out / 'marked',
          point_values={q: (points or {}).get(q, 1.0) for q in QS},
          key_data={'metadata': {'title': 'Practical 1', 'practical': 'yes'}})
    quiet(outputs.record, out, [path])
    return out, path


def _sheet(out):
    """The practical's gradebook tab, whatever its title."""
    import openpyxl
    wb = openpyxl.load_workbook(out / outputs.GRADEBOOK)
    return wb[next(n for n in wb.sheetnames
                   if n not in ('Exam stats', 'By question', 'Canvas scores'))]


def _stats(out):
    import openpyxl
    return openpyxl.load_workbook(out / outputs.GRADEBOOK)['Exam stats']


def _stat_row(out_ws, label) -> list:
    return next([c.value for c in row] for row in out_ws.iter_rows()
                if row[0].value == label)


def _stat(out_ws, label) -> list:
    """One Exam stats row's figures, past its label."""
    return [v for v in _stat_row(out_ws, label)[1:] if v is not None]


class PracticalPointsTest(unittest.TestCase):
    def setUp(self):
        self.out, self.csv = make_outputs(Path(tempfile.mkdtemp()))

    def tearDown(self):
        shutil.rmtree(self.out.parent)

    def graded(self):
        return outputs.load(self.csv)

    def results(self, *made):
        edits.save(self.out, list(made))
        return rm.open_folder(self.out)

    # ── the points on one sheet ────────────────────────────────────────────

    def test_the_key_carries_every_letter(self):
        self.assertEqual(self.graded().points_possible(), 8)

    def test_a_student_is_out_of_their_own_form(self):
        g = self.graded()
        self.assertEqual([g.points_possible(r) for r in g.students], [4, 4])
        self.assertEqual([g.total(r) for r in g.students], [2, 3])

    def test_the_forms_totals_are_reported_together(self):
        self.assertEqual(self.graded().student_possibles(), {4})
        self.assertEqual(outputs.all_possible([self.graded()]), {4})

    def test_a_dropped_question_leaves_the_sheet_it_was_on(self):
        g, _ = edits.apply([self.graded()],
                           [edits.make('credit', question='openQ_1A', action='drop')])
        self.assertEqual([g[0].points_possible(r) for r in g[0].students], [3, 4])

    # ── the analysis ───────────────────────────────────────────────────────

    def test_the_summary_says_every_student_is_out_of_four(self):
        self.assertEqual(rm.columns(self.results())[0].possible, 4)

    def test_a_student_row_carries_their_own_total(self):
        self.assertEqual([(s.name, s.total, s.possible) for s in rm.students(self.results())],
                         [('ABEL, X', 2, 4), ('BIRD, X', 3, 4)])

    def test_grade_bands_use_the_form_not_the_key(self):
        before, _ = rm.grade_counts(self.results())
        # 2/4 is 50%, a close F; 3/4 is 75%, a C. Against the key's 8 points
        # they would be 25% and 37.5%, both low F
        self.assertEqual({k: v for k, v in before.items() if v}, {'Close F': 1, 'C': 1})

    # ── the files ──────────────────────────────────────────────────────────

    def test_canvas_points_possible_is_what_a_student_could_earn(self):
        quiet(outputs.write, self.out, prefix='L')
        with (self.out / outputs.CANVAS).open() as f:
            rows = list(csv.reader(f))
        self.assertEqual(rows[1], ['    Points Possible', '', '4'])
        self.assertEqual([(r[1], r[2]) for r in rows[2:]], [('L1', '2'), ('L2', '3')])

    def test_the_gradebook_counts_one_form_not_the_key(self):
        quiet(outputs.write, self.out, prefix='L')
        ws = _stats(self.out)
        self.assertEqual(_stat(ws, 'Points possible'), [4, 4])
        # Every student is out of the same 4, so no Possible column is needed
        self.assertEqual([c.value for c in _sheet(self.out)[1]][-2:], ['Total', 'Adjusted'])


class UnevenFormsTest(unittest.TestCase):
    """Points given to one letter alone (the (2 pts) line the lab practical
    docs warn about) leave the forms out of different totals, so each
    student's own goes in a column of its own."""

    def setUp(self):
        self.out, self.csv = make_outputs(Path(tempfile.mkdtemp()),
                                          points={'openQ_1A': 2.0})

    def tearDown(self):
        shutil.rmtree(self.out.parent)

    def test_the_forms_are_out_of_different_totals(self):
        self.assertEqual(outputs.load(self.csv).student_possibles(), {4, 5})

    def test_the_summary_holds_no_single_total(self):
        self.assertIsNone(rm.columns(rm.open_folder(self.out))[0].possible)

    def test_the_gradebook_gives_each_student_their_own(self):
        quiet(outputs.write, self.out, prefix='L')
        ws = _sheet(self.out)
        self.assertEqual([c.value for c in ws[1]][-3:], ['Total', 'Adjusted', 'Possible'])
        self.assertEqual([r[-1] for r in ws.iter_rows(min_row=3, max_row=4,
                                                      values_only=True)], [5, 4])
        self.assertEqual(_stat(_stats(self.out), 'Points possible'), ['varies', 'varies'])

    def test_the_bands_count_each_student_against_their_own(self):
        quiet(outputs.write, self.out, prefix='L')
        band = _stat_row(_stats(self.out), 'B')[2]
        self.assertIn('SUMPRODUCT', band)
        self.assertIn('$V$3:$V$4', band)        # the Possible column


class WholeClassPointsTest(unittest.TestCase):
    """An exam every student answers in full keeps its single total and needs
    no Possible column."""

    def setUp(self):
        from test_edits import make_outputs as make_exam
        self.out, self.csvs = make_exam(Path(tempfile.mkdtemp()))

    def tearDown(self):
        shutil.rmtree(self.out.parent)

    def test_one_total_for_everyone(self):
        g = outputs.load(self.csvs[0])
        self.assertEqual(g.student_possibles(), {3})
        self.assertEqual([g.points_possible(r) for r in g.students], [3, 3, 3])

    def test_no_possible_column(self):
        import openpyxl
        quiet(outputs.write, self.out, prefix='L')
        wb = openpyxl.load_workbook(self.out / outputs.GRADEBOOK)
        header = [c.value for c in wb['Version A'][1]]
        self.assertEqual(header[-2:], ['Total', 'Adjusted'])


if __name__ == '__main__':
    unittest.main()
